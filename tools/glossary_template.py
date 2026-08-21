"""엑셀 용어집 양식 생성 — 계획서 §11 Phase 1 작업 1 (O-02: 엑셀).

    python -m tools.glossary_template [출력경로]

기본 출력: data/glossary_template.xlsx

**용어집을 채우기 전에 이 양식을 먼저 받아 갈 것.** 임의 배치로 500행을 채운
뒤에 맞추면 필수 필드(출처·신뢰도)가 빠져 다시 채워야 한다 — §5.1 이 그 둘을
"대역 충돌 시 판단하는 유일한 근거" 로 규정하기 때문이다.

개발망 전용이다. 폐쇄망에 반입하지 않는다 (§10.5).
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from tools.glossary_schema import (
    CONDITION_COLUMNS,
    LIST_SEP,
    SHEET_CONDITIONS,
    SHEET_GUIDE,
    SHEET_TERMS,
    TERM_COLUMNS,
    Column,
)

DEFAULT_OUT = Path("data/glossary_template.xlsx")

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_REQUIRED_FILL = PatternFill("solid", fgColor="C00000")
_HINT_FILL = PatternFill("solid", fgColor="EDEDED")
_EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")

#: 예시 행. 검수자가 감을 잡고 지우도록 색으로 구분한다.
_EXAMPLE_TERMS = [
    {
        "ko": "합동참모본부",
        "en": "Joint Chiefs of Staff",
        "en_abbr": "JCS",
        "ko_aliases": f"합참{LIST_SEP}합동참모부",
        "en_aliases": "ROK JCS",
        "domain_tag": f"편제{LIST_SEP}기관",
        "abbr_policy": "first_full_then_abbr",
        "source": "국방백서 국/영문판",
        "confidence": "verified",
        "note": "",
    },
    {
        "ko": "대령",
        "en": "Colonel",
        "en_abbr": "COL",
        "domain_tag": "계급",
        "abbr_policy": "always_full",
        "source": "군인사법",
        "confidence": "verified",
        "note": "군종에 따라 갈림 → '다의어' 시트에 함께 적을 것",
    },
]

_EXAMPLE_CONDITIONS = [
    {
        "ko": "대령",
        "field": "service",
        "value": f"육군{LIST_SEP}공군{LIST_SEP}해병대",
        "en": "Colonel",
    },
    {"ko": "대령", "field": "service", "value": "해군", "en": "Captain"},
]

GUIDE_LINES = [
    ("용어집 작성 안내", True),
    ("", False),
    ("■ 시트 구성", True),
    ("  · 용어    : 한 행에 용어 하나. 대부분 여기만 채우면 된다", False),
    ("  · 다의어  : 조건에 따라 대역이 갈리는 용어만 추가로 적는다", False),
    ("", False),
    ("■ 반드시 채워야 하는 열 (붉은 표시)", True),
    ("  한국어 / 영어 / 분류 / 출처 / 신뢰도", False),
    ("  출처와 신뢰도를 비우지 말 것. 나중에 같은 용어에 다른 대역이", False),
    ("  들어왔을 때 어느 쪽을 믿을지 판단하는 유일한 근거다.", False),
    ("", False),
    ("■ 비워도 되는 열", True),
    ("  ID / 우선순위 : 비우면 자동으로 채워진다", False),
    ("  영문약어 / 이형태 / 약어정책 / 메모 : 해당 없으면 비움", False),
    ("", False),
    ("■ 여러 값을 넣는 열", True),
    (f"  이형태와 분류는 세미콜론({LIST_SEP})으로 구분한다.", False),
    (f"  예: 합참{LIST_SEP}합동참모부", False),
    ("  쉼표는 쓰지 않는다 — 영문 용어 안에 쉼표가 들어갈 수 있다.", False),
    ("", False),
    ("■ 이형태에 적지 않아도 되는 것", True),
    ("  아래는 시스템이 자동으로 흡수하므로 적을 필요가 없다.", False),
    ("  · 띄어쓰기 차이   : '합동 참모 본부' = '합동참모본부'", False),
    ("  · 영문 대소문자   : 'JCS' = 'jcs'", False),
    ("  · 하이픈 종류     : 'ROK-US' = 'ROK–US'", False),
    ("  · 영문 복수형     : 'brigade' = 'brigades'", False),
    ("  · 전각/반각       : 'Ｋ２' = 'K2'", False),
    ("  적어야 하는 것은 '합참'처럼 형태가 다른 줄임말이다.", False),
    ("", False),
    ("■ 신뢰도", True),
    ("  verified  : 사람이 검수를 마쳤다. 확실할 때만", False),
    ("  probable  : 근거는 있으나 검수 전", False),
    ("  candidate : 후보. 검수 대기", False),
    ("", False),
    ("■ 약어정책", True),
    ("  first_full_then_abbr : 첫 등장은 'Joint Chiefs of Staff (JCS)',", False),
    ("                         이후는 'JCS'   ← 기관·부대에 주로 씀", False),
    ("  always_full          : 약어를 쓰지 않는다  ← 계급이 여기 해당", False),
    ("  always_abbr          : 항상 약어로만", False),
    ("  비움                 : 약어가 없거나 쓰지 않음", False),
    ("", False),
    ("■ 다의어 시트", True),
    ("  같은 한국어가 조건에 따라 다른 영어가 될 때만 쓴다.", False),
    ("  '한국어' 열은 용어 시트의 한국어와 글자까지 같아야 한다.", False),
    ("  조건분야는 비우면 service(군종)로 본다.", False),
    ("", False),
    ("  예) 대령 : 육군·공군·해병대는 Colonel, 해군은 Captain", False),
    ("      → 용어 시트에 '대령' 한 행 + 다의어 시트에 두 행", False),
    ("", False),
    ("■ 다 채운 뒤", True),
    ("  python -m tools.glossary_import <파일.xlsx>", False),
    ("  오류가 있으면 시트와 행 번호를 알려준다.", False),
]


def _write_header(ws, columns: tuple[Column, ...]) -> None:
    """1행 = 열 이름, 2행 = 설명."""
    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = col.width

        head = ws.cell(row=1, column=idx, value=col.header + (" *" if col.required else ""))
        head.font = Font(bold=True, color="FFFFFF", size=11)
        head.fill = _REQUIRED_FILL if col.required else _HEADER_FILL
        head.alignment = Alignment(horizontal="center", vertical="center")

        hint = ws.cell(row=2, column=idx, value=col.hint)
        hint.font = Font(size=9, color="595959")
        hint.fill = _HINT_FILL
        hint.alignment = Alignment(wrap_text=True, vertical="top")

    ws.row_dimensions[2].height = 34
    # 헤더와 설명은 항상 보이게 고정
    ws.freeze_panes = "A3"


def _add_validation(ws, columns: tuple[Column, ...], last_row: int = 1000) -> None:
    """드롭다운. 로더가 enum 으로 검증하므로 오타를 애초에 막는다 (§5.1)."""
    for idx, col in enumerate(columns, start=1):
        if not col.choices:
            continue
        letter = get_column_letter(idx)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(col.choices) + '"',
            allow_blank=not col.required,
            showDropDown=False,
        )
        dv.error = f"{col.header} 는 다음 중 하나여야 합니다: {', '.join(col.choices)}"
        dv.errorTitle = "허용되지 않는 값"
        ws.add_data_validation(dv)
        dv.add(f"{letter}3:{letter}{last_row}")


def _write_examples(ws, columns: tuple[Column, ...], rows: list[dict]) -> None:
    for offset, data in enumerate(rows):
        row = 3 + offset
        for idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=row, column=idx, value=data.get(col.key, ""))
            cell.fill = _EXAMPLE_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_template(path: Path) -> Path:
    wb = Workbook()

    terms = wb.active
    terms.title = SHEET_TERMS
    _write_header(terms, TERM_COLUMNS)
    _add_validation(terms, TERM_COLUMNS)
    _write_examples(terms, TERM_COLUMNS, _EXAMPLE_TERMS)

    conditions = wb.create_sheet(SHEET_CONDITIONS)
    _write_header(conditions, CONDITION_COLUMNS)
    _write_examples(conditions, CONDITION_COLUMNS, _EXAMPLE_CONDITIONS)

    guide = wb.create_sheet(SHEET_GUIDE)
    guide.column_dimensions["A"].width = 78
    for offset, (line, is_heading) in enumerate(GUIDE_LINES, start=1):
        cell = guide.cell(row=offset, column=1, value=line)
        cell.font = Font(bold=is_heading, size=12 if is_heading else 10)

    # 예시 행이 노란색이라는 것을 용어 시트에서도 알린다
    note_row = 3 + len(_EXAMPLE_TERMS) + 1
    marker = terms.cell(
        row=note_row,
        column=1,
        value="↑ 노란 행은 예시입니다. 지우고 채우세요. 작성 규칙은 '작성안내' 시트 참조.",
    )
    marker.font = Font(size=10, color="C00000", bold=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    args = argv if argv is not None else sys.argv[1:]
    out = Path(args[0]) if args else DEFAULT_OUT
    build_template(out)
    print(f"양식 생성: {out}")
    print(f"  시트: {SHEET_TERMS} / {SHEET_CONDITIONS} / {SHEET_GUIDE}")
    print("  채운 뒤: python -m tools.glossary_import <파일.xlsx>")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
