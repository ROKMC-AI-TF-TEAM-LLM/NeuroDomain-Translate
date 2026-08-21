"""엑셀 용어집 양식과 파서 — 계획서 §11 Phase 1 작업 1 (O-02: 엑셀).

가장 중요한 것은 **오류를 시트·행 번호로 알려주는가**다. 사람이 500행을
채우는 파일이라 "몇 번째 줄이 틀렸는지" 없이는 못 고친다 (R-11).

두 번째는 **틀린 항목을 통과시키지 않는가**다. 용어집은 번역을 강제하는
장치라 잘못된 항목 하나가 빠진 항목 백 개보다 나쁘다.
"""

from __future__ import annotations

from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl", reason="requirements-tools.txt 미설치")

from app.glossary.loader import load_glossary, load_meta  # noqa: E402
from tools.glossary_import import (  # noqa: E402
    FIRST_DATA_ROW,
    read_workbook,
    write_outputs,
)
from tools.glossary_schema import (  # noqa: E402
    CONDITION_COLUMNS,
    LIST_SEP,
    SHEET_CONDITIONS,
    SHEET_GUIDE,
    SHEET_TERMS,
    TERM_COLUMNS,
)
from tools.glossary_template import build_template  # noqa: E402


def make_workbook(
    tmp_path: Path,
    terms: list[dict],
    conditions: list[dict] | None = None,
    *,
    drop_columns: set[str] = frozenset(),
) -> Path:
    """테스트용 엑셀을 만든다. 양식과 같은 열 구조를 쓴다."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TERMS
    cols = [c for c in TERM_COLUMNS if c.key not in drop_columns]
    for idx, col in enumerate(cols, start=1):
        ws.cell(row=1, column=idx, value=col.header + (" *" if col.required else ""))
        ws.cell(row=2, column=idx, value=col.hint)
    for offset, data in enumerate(terms):
        for idx, col in enumerate(cols, start=1):
            ws.cell(row=FIRST_DATA_ROW + offset, column=idx, value=data.get(col.key, ""))

    cw = wb.create_sheet(SHEET_CONDITIONS)
    for idx, col in enumerate(CONDITION_COLUMNS, start=1):
        cw.cell(row=1, column=idx, value=col.header + (" *" if col.required else ""))
        cw.cell(row=2, column=idx, value=col.hint)
    for offset, data in enumerate(conditions or []):
        for idx, col in enumerate(CONDITION_COLUMNS, start=1):
            cw.cell(row=FIRST_DATA_ROW + offset, column=idx, value=data.get(col.key, ""))

    path = tmp_path / "glossary.xlsx"
    wb.save(path)
    return path


VALID = {
    "ko": "합동참모본부",
    "en": "Joint Chiefs of Staff",
    "en_abbr": "JCS",
    "ko_aliases": f"합참{LIST_SEP}합동참모부",
    "domain_tag": "편제",
    "abbr_policy": "first_full_then_abbr",
    "source": "국방백서",
    "confidence": "verified",
}


# ── 양식 ──────────────────────────────────────────────────────


def test_template_has_three_sheets(tmp_path: Path) -> None:
    from openpyxl import load_workbook as load

    path = build_template(tmp_path / "t.xlsx")
    wb = load(path)
    assert wb.sheetnames == [SHEET_TERMS, SHEET_CONDITIONS, SHEET_GUIDE]


def test_template_marks_required_columns(tmp_path: Path) -> None:
    """검수자가 무엇을 반드시 채워야 하는지 한눈에 보여야 한다."""
    from openpyxl import load_workbook as load

    wb = load(build_template(tmp_path / "t.xlsx"))
    headers = [c.value for c in wb[SHEET_TERMS][1]]
    required = {c.header for c in TERM_COLUMNS if c.required}
    for name in required:
        assert f"{name} *" in headers


def test_template_has_dropdowns_for_enums(tmp_path: Path) -> None:
    """로더가 enum 으로 검증하므로 오타 하나로 파일 전체가 거부된다 (§5.1)."""
    from openpyxl import load_workbook as load

    wb = load(build_template(tmp_path / "t.xlsx"))
    validations = wb[SHEET_TERMS].data_validations.dataValidation
    formulas = " ".join(v.formula1 for v in validations)
    assert "verified" in formulas
    assert "always_full" in formulas


def test_template_round_trips_through_the_importer(tmp_path: Path) -> None:
    """양식의 예시 행이 그대로 파싱돼야 예시가 맞다는 뜻이다."""
    report = read_workbook(build_template(tmp_path / "t.xlsx"))
    assert report.ok, [str(e) for e in report.errors]
    assert len(report.terms) == 2
    assert {t.ko for t in report.terms} == {"합동참모본부", "대령"}


# ── 파싱 ──────────────────────────────────────────────────────


def test_reads_a_valid_row(tmp_path: Path) -> None:
    report = read_workbook(make_workbook(tmp_path, [VALID]))
    assert report.ok, [str(e) for e in report.errors]
    term = report.terms[0]
    assert term.ko == "합동참모본부"
    assert term.en_abbr == "JCS"
    assert "합참" in term.ko_aliases


def test_ids_are_assigned_when_blank(tmp_path: Path) -> None:
    """검수자가 입력할 것을 줄인다. ID 는 사람이 정할 값이 아니다."""
    rows = [VALID, {**VALID, "ko": "국방부", "en": "Ministry of National Defense"}]
    report = read_workbook(make_workbook(tmp_path, rows))
    assert [t.id for t in report.terms] == ["T-0001", "T-0002"]
    assert report.assigned_ids == 2


def test_manual_ids_are_kept_and_not_reused(tmp_path: Path) -> None:
    """`id` 는 영구 불변이다 (§5.1). 이미 쓰인 번호를 다시 주면 안 된다."""
    rows = [
        {**VALID, "id": "T-0001"},
        {**VALID, "ko": "국방부", "en": "Ministry of National Defense"},
    ]
    report = read_workbook(make_workbook(tmp_path, rows))
    ids = [t.id for t in report.terms]
    assert ids[0] == "T-0001"
    assert ids[1] != "T-0001"
    assert len(set(ids)) == 2


def test_priority_is_derived_from_length(tmp_path: Path) -> None:
    """긴 표제어가 짧은 것을 이겨야 한다 (§6.3 겹침 해소)."""
    rows = [
        {**VALID, "ko": "제7기동군단", "en": "VII Maneuver Corps"},
        {**VALID, "ko": "군단", "en": "Corps"},
    ]
    report = read_workbook(make_workbook(tmp_path, rows))
    by_ko = {t.ko: t.priority for t in report.terms}
    assert by_ko["제7기동군단"] > by_ko["군단"]


def test_explicit_priority_wins(tmp_path: Path) -> None:
    report = read_workbook(make_workbook(tmp_path, [{**VALID, "priority": 42}]))
    assert report.terms[0].priority == 42


def test_corpus_freq_starts_at_zero(tmp_path: Path) -> None:
    """실측값은 Phase 1 작업 5 에서 채운다. 사람이 짐작할 값이 아니다."""
    report = read_workbook(make_workbook(tmp_path, [VALID]))
    assert report.terms[0].corpus_freq == 0


# ── 이형태 자동 생성 ──────────────────────────────────────────


def test_roman_numeral_alias_is_generated(tmp_path: Path) -> None:
    """`천무-Ⅱ` ↔ `천무-2`. 정규화로 흡수되지 않는 몇 안 되는 변형이다."""
    report = read_workbook(
        make_workbook(tmp_path, [{**VALID, "ko": "천무-Ⅱ", "en": "Chunmoo-II", "ko_aliases": ""}])
    )
    assert "천무-2" in report.terms[0].ko_aliases
    assert report.generated_aliases


def test_no_alias_for_variants_the_matcher_absorbs(tmp_path: Path) -> None:
    """띄어쓰기 · 대소문자 · 하이픈은 `normalize_key` 가 흡수한다 (§6.3).

    여기서 또 만들면 인덱스만 부풀고 얻는 것이 없다.
    """
    report = read_workbook(
        make_workbook(tmp_path, [{**VALID, "ko": "합동참모본부", "ko_aliases": ""}])
    )
    assert report.terms[0].ko_aliases == []
    assert report.generated_aliases == []


# ── 다의어 (§5.3) ─────────────────────────────────────────────


def test_conditions_are_merged_from_the_second_sheet(tmp_path: Path) -> None:
    path = make_workbook(
        tmp_path,
        [{**VALID, "ko": "대령", "en": "Colonel", "ko_aliases": "", "en_abbr": "COL"}],
        [
            {"ko": "대령", "value": f"육군{LIST_SEP}공군{LIST_SEP}해병대", "en": "Colonel"},
            {"ko": "대령", "value": "해군", "en": "Captain"},
        ],
    )
    report = read_workbook(path)
    assert report.ok, [str(e) for e in report.errors]

    term = report.terms[0]
    assert term.is_ambiguous
    assert term.conditions["field"] == "service"
    assert "Captain" in term.target_forms("ko2en")

    navy = [b for b in term.conditions["branches"] if "해군" in b["value"]]
    assert navy and navy[0]["en"] == "Captain"


def test_condition_for_unknown_term_is_an_error(tmp_path: Path) -> None:
    """표기가 어긋나면 조건이 조용히 사라진다. 그건 오역으로 이어진다."""
    path = make_workbook(
        tmp_path,
        [VALID],
        [{"ko": "대령", "value": "해군", "en": "Captain"}],
    )
    report = read_workbook(path)
    assert not report.ok
    assert any("대령" in str(e) for e in report.errors)


def test_single_branch_is_dropped_with_a_warning(tmp_path: Path) -> None:
    """분기가 하나면 다의어가 아니다. 다른 분기를 빠뜨렸을 수도 있다."""
    path = make_workbook(
        tmp_path,
        [{**VALID, "ko": "대령", "en": "Colonel", "ko_aliases": ""}],
        [{"ko": "대령", "value": "해군", "en": "Captain"}],
    )
    report = read_workbook(path)
    assert report.ok
    assert report.terms[0].conditions is None
    assert report.warnings


# ── 오류 보고 (R-11) ──────────────────────────────────────────


def test_missing_required_value_reports_the_row(tmp_path: Path) -> None:
    rows = [VALID, {**VALID, "ko": "국방부", "en": "", "source": "x"}]
    report = read_workbook(make_workbook(tmp_path, rows))
    assert not report.ok
    assert any(e.row == FIRST_DATA_ROW + 1 for e in report.errors)


def test_missing_source_is_rejected(tmp_path: Path) -> None:
    """출처는 대역 충돌 시 판단하는 유일한 근거다. 비우면 통과시키지 않는다."""
    report = read_workbook(make_workbook(tmp_path, [{**VALID, "source": ""}]))
    assert not report.ok
    assert any("source" in str(e) for e in report.errors)


def test_bad_confidence_is_rejected_with_the_row(tmp_path: Path) -> None:
    report = read_workbook(make_workbook(tmp_path, [{**VALID, "confidence": "아마도"}]))
    assert not report.ok
    assert any(e.row == FIRST_DATA_ROW and "confidence" in str(e) for e in report.errors)


def test_duplicate_korean_is_rejected(tmp_path: Path) -> None:
    """같은 한국어가 두 번 나오면 어느 대역이 맞는지 알 수 없다."""
    report = read_workbook(make_workbook(tmp_path, [VALID, VALID]))
    assert not report.ok
    assert any("중복" in str(e) for e in report.errors)


def test_missing_required_column_is_reported(tmp_path: Path) -> None:
    report = read_workbook(make_workbook(tmp_path, [VALID], drop_columns={"source"}))
    assert not report.ok
    assert any("필수 열" in str(e) for e in report.errors)


def test_comma_separator_gets_a_warning(tmp_path: Path) -> None:
    """쉼표는 영문 용어 안에 들어갈 수 있어 구분자로 쓰지 않는다."""
    report = read_workbook(make_workbook(tmp_path, [{**VALID, "ko_aliases": "합참,합동참모부"}]))
    assert report.warnings


def test_blank_rows_are_skipped(tmp_path: Path) -> None:
    report = read_workbook(make_workbook(tmp_path, [VALID, {}, {}]))
    assert report.ok
    assert len(report.terms) == 1


# ── 출력 ──────────────────────────────────────────────────────


def test_output_loads_with_the_app_loader(tmp_path: Path) -> None:
    """**이것이 최종 확인이다.** 파서가 만든 파일을 앱이 읽을 수 있어야 한다."""
    path = make_workbook(
        tmp_path,
        [
            VALID,
            {**VALID, "ko": "대령", "en": "Colonel", "ko_aliases": "", "en_abbr": "COL"},
        ],
        [
            {"ko": "대령", "value": f"육군{LIST_SEP}공군", "en": "Colonel"},
            {"ko": "대령", "value": "해군", "en": "Captain"},
        ],
    )
    report = read_workbook(path)
    assert report.ok, [str(e) for e in report.errors]

    out = tmp_path / "glossary.jsonl"
    meta = tmp_path / "glossary.meta.json"
    write_outputs(report, out, meta)

    terms = load_glossary(out)
    assert len(terms) == 2
    assert load_meta(meta).count == 2


def test_meta_version_increments(tmp_path: Path) -> None:
    """인덱스 재빌드 판단에 쓴다. 갱신할 때마다 올라가야 한다 (§6.7)."""
    report = read_workbook(make_workbook(tmp_path, [VALID]))
    out, meta = tmp_path / "g.jsonl", tmp_path / "g.meta.json"

    write_outputs(report, out, meta)
    assert load_meta(meta).version == 1
    write_outputs(report, out, meta)
    assert load_meta(meta).version == 2
